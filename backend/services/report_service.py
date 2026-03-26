"""Reporting module business logic — aggregation for all 5 standard reports."""

from __future__ import annotations

import io
from datetime import date

from sqlalchemy import distinct, func, case
from sqlalchemy.orm import Session

from config import DEMO_DATE
from models.capacity import Allocation
from models.financial import Actuals, Baseline, Forecast
from models.organization import CostCenter, LineOfBusiness
from models.people import Person, RateTable, RoleType
from models.projects import Project
from models.financial import ExternalCostType
from models.reporting import ForecastSnapshot
from schemas.common import CurrentUser
from services.calculations import add_months, compute_plan_drift


# ---------------------------------------------------------------------------
# Role-scoping helper
# ---------------------------------------------------------------------------

def _get_scoped_project_ids(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
) -> list[str]:
    """Return project IDs visible to the user, optionally filtered."""
    filters = filters or {}

    q = db.query(Project.id).filter(Project.is_active.is_(True))

    # Role scoping
    if user.role == "project_lead":
        q = q.filter(Project.id.in_(user.project_ids))
    elif user.role == "cost_center_owner" and user.cost_center_id:
        # Projects that have allocations from this user's cost center
        cc_person_ids = [
            r[0] for r in db.query(Person.id)
            .filter(Person.cost_center_id == user.cost_center_id)
            .all()
        ]
        if cc_person_ids:
            project_ids_via_alloc = [
                r[0] for r in db.query(distinct(Allocation.project_id))
                .filter(Allocation.person_id.in_(cc_person_ids))
                .all()
            ]
            q = q.filter(Project.id.in_(project_ids_via_alloc))
        else:
            return []
    # controller / executive see everything

    # Apply common filters
    if filters.get("lob"):
        q = q.filter(Project.lob_id == filters["lob"])
    if filters.get("status"):
        q = q.filter(Project.status == filters["status"])
    if filters.get("rag"):
        q = q.filter(Project.rag_status == filters["rag"])
    if filters.get("type"):
        if filters["type"] == "service":
            q = q.filter(Project.is_service.is_(True))
        elif filters["type"] == "project":
            q = q.filter(Project.is_service.is_(False))

    return [r[0] for r in q.all()]


def _apply_project_filter(q, project_ids: list[str], project_id_col):
    """Apply project ID IN filter to a query."""
    if project_ids:
        return q.filter(project_id_col.in_(project_ids))
    return q.filter(project_id_col == "__none__")  # empty result


# ---------------------------------------------------------------------------
# 1. Programme / Multi-Project Rollup
# ---------------------------------------------------------------------------

def compute_programme_rollup(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
    grouping: str = "lob",
    fiscal_year: int | None = None,
) -> dict:
    """Consolidated financials across multiple projects."""
    filters = filters or {}
    # RPT-03: custom project group overrides normal scoping
    if filters.get("project_ids"):
        project_ids = filters["project_ids"]
    else:
        project_ids = _get_scoped_project_ids(db, user, filters)

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": [], "total": 0}

    # Load projects with financials
    projects = db.query(Project).filter(Project.id.in_(project_ids)).all()
    lobs = {l.id: l.name for l in db.query(LineOfBusiness).all()}

    # Build optional year filter (None = lifetime totals)
    year_prefix = str(fiscal_year) if fiscal_year else None

    rows = []
    total_baseline = 0.0
    total_forecast = 0.0
    total_actuals = 0.0

    for p in projects:
        bl_q = db.query(func.coalesce(func.sum(Baseline.amount_eur), 0)).filter(Baseline.project_id == p.id)
        fc_q = db.query(func.coalesce(func.sum(Forecast.amount_eur), 0)).filter(Forecast.project_id == p.id)
        ac_q = db.query(func.coalesce(func.sum(Actuals.amount_eur), 0)).filter(Actuals.project_id == p.id, Actuals.month <= DEMO_DATE)
        if year_prefix:
            bl_q = bl_q.filter(func.substr(Baseline.month, 1, 4) == year_prefix)
            fc_q = fc_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
            ac_q = ac_q.filter(func.substr(Actuals.month, 1, 4) == year_prefix)
        baseline = float(bl_q.scalar())
        forecast = float(fc_q.scalar())
        actuals = float(ac_q.scalar())
        remaining = forecast - actuals
        variance = forecast - baseline
        variance_pct = compute_plan_drift(forecast, baseline)

        total_baseline += baseline
        total_forecast += forecast
        total_actuals += actuals

        rows.append({
            "project_id": p.id,
            "project_name": p.name,
            "lob_id": p.lob_id,
            "lob_name": lobs.get(p.lob_id, ""),
            "status": p.status,
            "rag": p.rag_status,
            "baseline_budget": round(baseline, 2),
            "current_forecast": round(forecast, 2),
            "actuals_to_date": round(actuals, 2),
            "remaining_forecast": round(max(remaining, 0), 2),
            "variance": round(variance, 2),
            "variance_pct": round(variance_pct, 1),
        })

    # Sort rows by variance descending (biggest overrun first)
    rows.sort(key=lambda r: r["variance"], reverse=True)

    # KPIs
    overall_variance = total_forecast - total_baseline
    overall_variance_pct = compute_plan_drift(total_forecast, total_baseline)
    rag_counts = {"green": 0, "amber": 0, "red": 0}
    for r in rows:
        if r["rag"] in rag_counts:
            rag_counts[r["rag"]] += 1

    kpis = {
        "project_count": len(rows),
        "total_baseline": round(total_baseline, 2),
        "total_forecast": round(total_forecast, 2),
        "overall_variance": round(overall_variance, 2),
        "overall_variance_pct": round(overall_variance_pct, 1),
        "rag_distribution": rag_counts,
    }

    # Chart data: per-project baseline/forecast/actuals
    chart_data = [
        {
            "name": r["project_name"][:20],
            "baseline": r["baseline_budget"],
            "forecast": r["current_forecast"],
            "actuals": r["actuals_to_date"],
        }
        for r in rows[:15]  # limit chart to 15 projects
    ]

    return {"kpis": kpis, "rows": rows, "chart_data": chart_data, "total": len(rows)}


# ---------------------------------------------------------------------------
# 2. Cost Center Financial Summary
# ---------------------------------------------------------------------------

def compute_cc_financial_summary(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
    fiscal_year: int | None = None,
) -> dict:
    """Financial picture for a cost center across all contributing projects."""
    filters = filters or {}
    project_ids = _get_scoped_project_ids(db, user, filters)
    year_prefix = str(fiscal_year) if fiscal_year else None

    # Determine cost center(s)
    target_cc = filters.get("cost_center")
    if not target_cc and user.role == "cost_center_owner":
        target_cc = user.cost_center_id

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": {}, "total": 0}

    # Get people in the target CC (for internal cost attribution)
    if target_cc:
        cc_person_ids = [
            r[0] for r in db.query(Person.id)
            .filter(Person.cost_center_id == target_cc)
            .all()
        ]
        # Filter projects to those with allocations from this CC
        project_ids_via_alloc = [
            r[0] for r in db.query(distinct(Allocation.project_id))
            .filter(Allocation.person_id.in_(cc_person_ids))
            .all()
        ] if cc_person_ids else []
        project_ids = [pid for pid in project_ids if pid in project_ids_via_alloc]
    else:
        cc_person_ids = None

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": {}, "total": 0}

    projects = db.query(Project).filter(Project.id.in_(project_ids)).all()

    # Build project → cost center name lookup (primary CC = CC with most allocations)
    cc_name_map: dict[str, str] = {}
    if target_cc:
        cc_obj = db.query(CostCenter).filter(CostCenter.id == target_cc).first()
        cc_label = cc_obj.name if cc_obj else target_cc
        for pid in project_ids:
            cc_name_map[pid] = cc_label
    else:
        for pid in project_ids:
            top_cc = (
                db.query(CostCenter.name, func.count(Allocation.id).label("cnt"))
                .join(Person, Person.id == Allocation.person_id)
                .join(CostCenter, CostCenter.id == Person.cost_center_id)
                .filter(Allocation.project_id == pid)
                .group_by(CostCenter.name)
                .order_by(func.count(Allocation.id).desc())
                .first()
            )
            cc_name_map[pid] = top_cc[0] if top_cc else "—"

    rows = []
    total_internal_hours = 0.0
    total_internal_cost = 0.0
    total_external_cost = 0.0

    for p in projects:
        # Internal: forecast rows where category=internal
        int_q = db.query(
            func.coalesce(func.sum(Forecast.hours), 0),
            func.coalesce(func.sum(Forecast.amount_eur), 0),
        ).filter(Forecast.project_id == p.id, Forecast.category == "internal")
        if year_prefix:
            int_q = int_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
        int_hours, int_cost = int_q.one()

        # External: forecast rows where category=external
        ext_q = db.query(func.coalesce(func.sum(Forecast.amount_eur), 0)).filter(
            Forecast.project_id == p.id, Forecast.category == "external")
        if year_prefix:
            ext_q = ext_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
        ext_cost = float(ext_q.scalar())

        int_hours = float(int_hours)
        int_cost = float(int_cost)
        total_cost = int_cost + ext_cost
        total_internal_hours += int_hours
        total_internal_cost += int_cost
        total_external_cost += ext_cost

        rows.append({
            "project_id": p.id,
            "cost_center_name": cc_name_map.get(p.id, "—"),
            "project_name": p.name,
            "internal_hours": round(int_hours, 1),
            "internal_cost": round(int_cost, 2),
            "external_cost": round(ext_cost, 2),
            "total_cost": round(total_cost, 2),
            "pct_of_cc_budget": 0.0,  # calculated after totals
            "status": p.status,
        })

    # Calculate percentages
    grand_total = total_internal_cost + total_external_cost
    for r in rows:
        r["pct_of_cc_budget"] = round(
            (r["total_cost"] / grand_total * 100) if grand_total > 0 else 0, 1
        )

    # Total actuals for KPI
    total_actuals = float(
        db.query(func.coalesce(func.sum(Actuals.amount_eur), 0))
        .filter(Actuals.project_id.in_(project_ids), Actuals.month <= DEMO_DATE, func.substr(Actuals.month, 1, 4) == year_prefix)
        .scalar()
    )

    kpis = {
        "total_budget_allocated": round(grand_total, 2),
        "total_actuals": round(total_actuals, 2),
        "total_internal_cost": round(total_internal_cost, 2),
        "total_external_cost": round(total_external_cost, 2),
        "active_project_count": len(rows),
    }

    # Chart data: pie (per-project share) + cost type breakdown
    pie_data = [
        {"name": r["project_name"][:20], "value": r["total_cost"]}
        for r in rows if r["total_cost"] > 0
    ]

    # Spend per cost type (internal vs external)
    cost_type_data = [
        {"name": "Internal", "value": round(total_internal_cost, 2)},
        {"name": "External", "value": round(total_external_cost, 2)},
    ]

    # Internal spend breakdown by role type
    int_breakdown_q = (
        db.query(Forecast.sub_category, func.sum(Forecast.amount_eur))
        .filter(Forecast.project_id.in_(project_ids), Forecast.category == "internal")
    )
    if year_prefix:
        int_breakdown_q = int_breakdown_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
    int_breakdown_q = int_breakdown_q.group_by(Forecast.sub_category)
    role_names = {r.id: r.name for r in db.query(RoleType).all()}
    internal_breakdown = [
        {"name": role_names.get(sc, sc), "value": round(float(amt), 2)}
        for sc, amt in int_breakdown_q.all() if float(amt) > 0
    ]
    internal_breakdown.sort(key=lambda x: x["value"], reverse=True)

    # External spend breakdown by cost type
    ext_breakdown_q = (
        db.query(Forecast.sub_category, func.sum(Forecast.amount_eur))
        .filter(Forecast.project_id.in_(project_ids), Forecast.category == "external")
    )
    if year_prefix:
        ext_breakdown_q = ext_breakdown_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
    ext_breakdown_q = ext_breakdown_q.group_by(Forecast.sub_category)
    cost_type_names = {c.id: c.name for c in db.query(ExternalCostType).all()}
    external_breakdown = [
        {"name": cost_type_names.get(sc, sc), "value": round(float(amt), 2)}
        for sc, amt in ext_breakdown_q.all() if float(amt) > 0
    ]
    external_breakdown.sort(key=lambda x: x["value"], reverse=True)

    return {
        "kpis": kpis,
        "rows": rows,
        "chart_data": {
            "pie": pie_data,
            "cost_type": cost_type_data,
            "internal_breakdown": internal_breakdown,
            "external_breakdown": external_breakdown,
        },
        "total": len(rows),
    }


# ---------------------------------------------------------------------------
# 3. Vendor Spend Analysis
# ---------------------------------------------------------------------------

def compute_vendor_spend(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
    fiscal_year: int | None = None,
) -> dict:
    """External spending analysis by vendor across the portfolio."""
    filters = filters or {}
    project_ids = _get_scoped_project_ids(db, user, filters)
    year_prefix = str(fiscal_year) if fiscal_year else None

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": {}, "total": 0}

    base_q = db.query(Forecast).filter(
        Forecast.project_id.in_(project_ids),
        Forecast.vendor.isnot(None),
        Forecast.vendor != "",
    )
    if year_prefix:
        base_q = base_q.filter(func.substr(Forecast.month, 1, 4) == year_prefix)
    if filters.get("vendor"):
        base_q = base_q.filter(Forecast.vendor == filters["vendor"])
    # RPT-05: expense cost type filter
    if filters.get("expense_cost_type"):
        base_q = base_q.filter(Forecast.sub_category == filters["expense_cost_type"])

    forecast_rows = base_q.all()

    # Group by vendor
    vendor_map: dict[str, dict] = {}
    for f in forecast_rows:
        v = f.vendor
        if v not in vendor_map:
            vendor_map[v] = {
                "vendor_name": v,
                "total_ordered": 0.0,
                "total_invoiced": 0.0,
                "total_open": 0.0,
                "total_accruals": 0.0,
                "projects": set(),
                "po_numbers": set(),
                "cost_types": {},  # RPT-05: track cost types
            }
        amt = float(f.amount_eur or 0)
        status = f.ext_status or "planned"

        if status in ("invoiced",):
            vendor_map[v]["total_invoiced"] += amt
        elif status in ("ordered", "goods_received"):
            vendor_map[v]["total_ordered"] += amt
        elif status in ("accrual",):
            vendor_map[v]["total_accruals"] += amt
        else:
            vendor_map[v]["total_open"] += amt

        vendor_map[v]["projects"].add(f.project_id)
        if f.po_number:
            vendor_map[v]["po_numbers"].add(f.po_number)
        # RPT-05: track cost type frequency
        ct = f.sub_category or ""
        if ct:
            vendor_map[v]["cost_types"][ct] = vendor_map[v]["cost_types"].get(ct, 0) + 1

    rows = []
    total_spend = 0.0
    total_open = 0.0
    for v_data in vendor_map.values():
        vendor_total = (
            v_data["total_ordered"] + v_data["total_invoiced"]
            + v_data["total_open"] + v_data["total_accruals"]
        )
        total_spend += vendor_total
        total_open += v_data["total_open"] + v_data["total_ordered"]
        # RPT-05: dominant cost type for vendor
        ct_map = v_data["cost_types"]
        dominant_ct = max(ct_map, key=ct_map.get) if ct_map else ""
        rows.append({
            "vendor_name": v_data["vendor_name"],
            "expense_cost_type": dominant_ct,
            "total_ordered": round(v_data["total_ordered"], 2),
            "total_invoiced": round(v_data["total_invoiced"], 2),
            "total_open": round(v_data["total_open"], 2),
            "total_accruals": round(v_data["total_accruals"], 2),
            "project_count": len(v_data["projects"]),
            "po_count": len(v_data["po_numbers"]),
        })

    rows.sort(key=lambda r: (
        r["total_ordered"] + r["total_invoiced"] + r["total_open"] + r["total_accruals"]
    ), reverse=True)

    kpis = {
        "total_vendor_spend": round(total_spend, 2),
        "active_vendor_count": len(rows),
        "total_po_count": sum(r["po_count"] for r in rows),
        "open_commitments": round(total_open, 2),
    }

    # Chart: top 10 horizontal bar
    bar_data = [
        {
            "vendor": r["vendor_name"][:25],
            "total": round(
                r["total_ordered"] + r["total_invoiced"]
                + r["total_open"] + r["total_accruals"], 2
            ),
        }
        for r in rows[:10]
    ]

    return {
        "kpis": kpis,
        "rows": rows,
        "chart_data": {"bar": bar_data},
        "total": len(rows),
    }


def compute_vendor_drilldown(
    db: Session,
    user: CurrentUser,
    vendor_name: str,
) -> list[dict]:
    """Individual line items for a specific vendor."""
    project_ids = _get_scoped_project_ids(db, user)

    if not project_ids:
        return []

    rows = (
        db.query(Forecast)
        .filter(
            Forecast.project_id.in_(project_ids),
            Forecast.vendor == vendor_name,
        )
        .order_by(Forecast.project_id, Forecast.month)
        .all()
    )

    # Get project names
    proj_names = {
        p.id: p.name
        for p in db.query(Project).filter(
            Project.id.in_(list({r.project_id for r in rows}))
        ).all()
    }

    return [
        {
            "project_id": r.project_id,
            "project_name": proj_names.get(r.project_id, ""),
            "month": r.month,
            "cost_type": r.sub_category,
            "amount": round(float(r.amount_eur or 0), 2),
            "status": r.ext_status,
            "po_number": r.po_number,
        }
        for r in rows
    ]


# ---------------------------------------------------------------------------
# 4. Forecast Accuracy
# ---------------------------------------------------------------------------

def compute_forecast_accuracy(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
    horizon_months: int = 6,
    fiscal_year: int | None = None,
) -> dict:
    """Compare historical forecast snapshots against actuals."""
    filters = filters or {}
    project_ids = _get_scoped_project_ids(db, user, filters)
    year_prefix = str(fiscal_year) if fiscal_year else None

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": [], "total": 0}

    snapshot_month = add_months(DEMO_DATE, -horizon_months)
    lobs = {l.id: l.name for l in db.query(LineOfBusiness).all()}

    # Get snapshot forecasts closest to the target month
    snapshots = (
        db.query(ForecastSnapshot)
        .filter(
            ForecastSnapshot.project_id.in_(project_ids),
            ForecastSnapshot.snapshot_month <= snapshot_month,
        )
        .order_by(ForecastSnapshot.project_id, ForecastSnapshot.snapshot_month.desc())
        .all()
    )

    # Keep only the latest snapshot per project that is <= snapshot_month
    latest_snapshots: dict[str, float] = {}
    for s in snapshots:
        if s.project_id not in latest_snapshots:
            latest_snapshots[s.project_id] = float(s.forecast_total)

    projects = db.query(Project).filter(Project.id.in_(project_ids)).all()

    rows = []
    total_accuracy = 0.0
    within_5 = 0
    above_15 = 0
    net_bias = 0.0  # positive = over-forecast, negative = under-forecast

    for p in projects:
        forecast_val = latest_snapshots.get(p.id)
        if forecast_val is None:
            continue

        ac_q = db.query(func.coalesce(func.sum(Actuals.amount_eur), 0)).filter(
            Actuals.project_id == p.id, Actuals.month <= DEMO_DATE)
        if year_prefix:
            ac_q = ac_q.filter(func.substr(Actuals.month, 1, 4) == year_prefix)
        actual_val = float(ac_q.scalar())

        if actual_val == 0:
            continue

        variance = forecast_val - actual_val
        variance_pct = abs(variance / actual_val * 100) if actual_val else 0
        accuracy = 100 - variance_pct

        if variance_pct < 5:
            rating = "green"
            within_5 += 1
        elif variance_pct <= 15:
            rating = "amber"
        else:
            rating = "red"
            above_15 += 1

        total_accuracy += max(accuracy, 0)
        net_bias += variance  # positive = over-forecast

        rows.append({
            "project_id": p.id,
            "project_name": p.name,
            "lob_id": p.lob_id,
            "lob_name": lobs.get(p.lob_id, ""),
            "forecast_value": round(forecast_val, 2),
            "actual_value": round(actual_val, 2),
            "variance": round(variance, 2),
            "variance_pct": round(variance_pct, 1),
            "accuracy_rating": rating,
        })

    avg_accuracy = round(total_accuracy / len(rows), 1) if rows else 0
    bias_direction = "over" if net_bias > 0 else "under" if net_bias < 0 else "neutral"

    kpis = {
        "avg_accuracy_pct": avg_accuracy,
        "within_5_count": within_5,
        "above_15_count": above_15,
        "bias_direction": bias_direction,
        "bias_amount": round(abs(net_bias), 2),
    }

    # Scatter chart data
    chart_data = [
        {
            "project_name": r["project_name"][:15],
            "forecast": r["forecast_value"],
            "actual": r["actual_value"],
            "lob": r["lob_name"],
            "rating": r["accuracy_rating"],
        }
        for r in rows
    ]

    return {"kpis": kpis, "rows": rows, "chart_data": chart_data, "total": len(rows)}


# ---------------------------------------------------------------------------
# 5. Year-over-Year Comparison
# ---------------------------------------------------------------------------

def compute_year_over_year(
    db: Session,
    user: CurrentUser,
    filters: dict | None = None,
    fy_current: int = 2026,
    fy_previous: int = 2025,
    show_monthly: bool = False,
    months_filter: list[int] | None = None,
) -> dict:
    """Compare portfolio spending between fiscal years.

    RPT-07/08: supports annual (default) and monthly modes with LoB/Project columns.
    """
    filters = filters or {}
    project_ids = _get_scoped_project_ids(db, user, filters)

    if not project_ids:
        return {"kpis": {}, "rows": [], "chart_data": [], "total": 0}

    cost_type = filters.get("cost_type")

    month_names = [
        "Jan", "Feb", "Mar", "Apr", "May", "Jun",
        "Jul", "Aug", "Sep", "Oct", "Nov", "Dec",
    ]

    # Build project info lookup (RPT-07: LoB + Project columns)
    project_info: dict[str, dict] = {}
    for p in db.query(Project).filter(Project.id.in_(project_ids)).all():
        lob = db.query(LineOfBusiness).filter(LineOfBusiness.id == p.lob_id).first()
        project_info[p.id] = {
            "project_name": p.name,
            "lob_id": p.lob_id or "",
            "lob_name": lob.name if lob else "",
        }

    def _base_query(year: int):
        q = db.query(
            Actuals.project_id,
            func.substr(Actuals.month, 6, 2),  # MM
            func.sum(Actuals.amount_eur),
        ).filter(
            Actuals.project_id.in_(project_ids),
            func.substr(Actuals.month, 1, 4) == str(year),
        )
        if cost_type and cost_type != "all":
            if cost_type == "internal":
                q = q.filter(Actuals.category == "internal")
            elif cost_type == "external":
                q = q.filter(Actuals.category == "external")
        return q

    # Fetch per-project, per-month actuals for both years
    def _get_spend_by_project_month(year: int) -> dict[str, dict[int, float]]:
        q = _base_query(year).group_by(Actuals.project_id, func.substr(Actuals.month, 6, 2))
        result: dict[str, dict[int, float]] = {}
        for pid, mm, amt in q.all():
            result.setdefault(pid, {})[int(mm)] = round(float(amt), 2)
        return result

    current_by_pm = _get_spend_by_project_month(fy_current)
    previous_by_pm = _get_spend_by_project_month(fy_previous)

    all_pids = sorted(set(current_by_pm.keys()) | set(previous_by_pm.keys()))

    # KPI computation (always uses all months for YTD)
    demo_month = int(DEMO_DATE.split("-")[1])
    ytd_current = 0.0
    ytd_previous = 0.0
    for pid in all_pids:
        for m in range(1, demo_month + 1):
            ytd_current += current_by_pm.get(pid, {}).get(m, 0.0)
            ytd_previous += previous_by_pm.get(pid, {}).get(m, 0.0)
    ytd_delta = ytd_current - ytd_previous

    kpis = {
        "fy_current_ytd": round(ytd_current, 2),
        "fy_previous_ytd": round(ytd_previous, 2),
        "ytd_delta": round(ytd_delta, 2),
        "trajectory": "higher" if ytd_delta > 0 else "lower" if ytd_delta < 0 else "flat",
        "fy_current_label": str(fy_current),
        "fy_previous_label": str(fy_previous),
    }

    rows = []

    if show_monthly:
        # RPT-08: Monthly detail mode — one row per project per month
        active_months = months_filter if months_filter else list(range(1, 13))
        # Cumulative tracking per project
        cum_current: dict[str, float] = {pid: 0.0 for pid in all_pids}
        cum_previous: dict[str, float] = {pid: 0.0 for pid in all_pids}

        for month_num in range(1, 13):
            for pid in all_pids:
                c = current_by_pm.get(pid, {}).get(month_num, 0.0)
                p = previous_by_pm.get(pid, {}).get(month_num, 0.0)
                cum_current[pid] += c
                cum_previous[pid] += p

                if month_num not in active_months:
                    continue

                delta = c - p
                delta_pct = round((delta / p * 100), 1) if p > 0 else 0.0
                info = project_info.get(pid, {})

                rows.append({
                    "month": month_names[month_num - 1],
                    "month_num": month_num,
                    "project_id": pid,
                    "project_name": info.get("project_name", pid),
                    "lob_id": info.get("lob_id", ""),
                    "lob_name": info.get("lob_name", ""),
                    "fy_current": round(c, 2),
                    "fy_previous": round(p, 2),
                    "delta": round(delta, 2),
                    "delta_pct": delta_pct,
                    "cumulative_current": round(cum_current[pid], 2),
                    "cumulative_previous": round(cum_previous[pid], 2),
                })
    else:
        # RPT-07/08: Annual mode — one row per project
        for pid in all_pids:
            c = sum(current_by_pm.get(pid, {}).values())
            p = sum(previous_by_pm.get(pid, {}).values())
            delta = c - p
            delta_pct = round((delta / p * 100), 1) if p > 0 else 0.0
            info = project_info.get(pid, {})

            rows.append({
                "project_id": pid,
                "project_name": info.get("project_name", pid),
                "lob_id": info.get("lob_id", ""),
                "lob_name": info.get("lob_name", ""),
                "fy_current": round(c, 2),
                "fy_previous": round(p, 2),
                "delta": round(delta, 2),
                "delta_pct": delta_pct,
            })

    # Chart data: aggregate across all projects by month (backward compat)
    chart_data = []
    agg_cum_c = 0.0
    agg_cum_p = 0.0
    for month_num in range(1, 13):
        mc = sum(current_by_pm.get(pid, {}).get(month_num, 0.0) for pid in all_pids)
        mp = sum(previous_by_pm.get(pid, {}).get(month_num, 0.0) for pid in all_pids)
        agg_cum_c += mc
        agg_cum_p += mp
        chart_data.append({
            "month": month_names[month_num - 1],
            "fy_current": round(agg_cum_c, 2),
            "fy_previous": round(agg_cum_p, 2),
            "fy_current_monthly": round(mc, 2),
            "fy_previous_monthly": round(mp, 2),
        })

    return {"kpis": kpis, "rows": rows, "chart_data": chart_data, "total": len(rows)}


# ---------------------------------------------------------------------------
# Excel Export
# ---------------------------------------------------------------------------

def generate_excel_export(
    report_name: str,
    headers: list[str],
    data_rows: list[list],
    summary_row: list | None = None,
    active_filters: dict | None = None,
) -> io.BytesIO:
    """Generate an Excel workbook in-memory and return as BytesIO."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Report"

    # Header row (bold, blue background)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill

    # Data rows
    for row_idx, data_row in enumerate(data_rows, 2):
        for col_idx, value in enumerate(data_row, 1):
            ws.cell(row=row_idx, column=col_idx, value=value)

    # Summary row (bold)
    if summary_row:
        summary_row_idx = len(data_rows) + 2
        summary_font = Font(bold=True)
        for col_idx, value in enumerate(summary_row, 1):
            cell = ws.cell(row=summary_row_idx, column=col_idx, value=value)
            cell.font = summary_font

    # Auto-width columns
    for col in ws.columns:
        max_length = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 30)

    # Metadata sheet
    meta_ws = wb.create_sheet("Metadata")
    meta_ws.cell(row=1, column=1, value="Report Name").font = Font(bold=True)
    meta_ws.cell(row=1, column=2, value=report_name)
    meta_ws.cell(row=2, column=1, value="Export Date").font = Font(bold=True)
    meta_ws.cell(row=2, column=2, value=date.today().isoformat())
    meta_ws.cell(row=3, column=1, value="Active Filters").font = Font(bold=True)
    if active_filters:
        filter_str = ", ".join(f"{k}={v}" for k, v in active_filters.items() if v)
        meta_ws.cell(row=3, column=2, value=filter_str or "None")
    else:
        meta_ws.cell(row=3, column=2, value="None")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
