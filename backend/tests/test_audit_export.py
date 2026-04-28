"""Tests for services/audit_export.py and the export router endpoint."""

from __future__ import annotations

from datetime import datetime, timedelta

import pytest

from models.system import AuditLog
from services.audit_export import CSV_COLUMNS, export_csv, export_xlsx
from services.audit_query import query_audit_log


HEADERS_CTRL = {"X-Current-User": "persona-controller"}


@pytest.fixture
def some_rows(db, seed_personas):
    base = datetime(2026, 4, 20, 9, 0, 0)
    for i in range(3):
        db.add(AuditLog(
            timestamp=base + timedelta(hours=i),
            user_person_id="p-dev-1",
            entity_type="cost_center",
            entity_id=f"cc-{i}",
            entity_name=f"CC {i}",
            action="create" if i == 0 else "update",
            field_changed=None if i == 0 else "name",
            old_value=None,
            new_value=f"New {i}",
            category="master_data",
        ))
    db.commit()
    items, _ = query_audit_log(db)
    return items


class TestExportCSV:
    def test_writes_bom_and_header(self, some_rows):
        buf = export_csv(some_rows, exported_by="Anna", filter_summary="entity_type=cost_center")
        data = buf.getvalue()
        # UTF-8 BOM at start
        assert data.startswith(b"\xef\xbb\xbf")
        # Decoded text contains preamble + header columns
        text = data.decode("utf-8-sig")
        assert "# CRETA Audit Log Export" in text
        assert "Anna" in text
        assert "entity_type=cost_center" in text
        # Each column heading appears
        for col in CSV_COLUMNS:
            assert col in text

    def test_writes_data_rows(self, some_rows):
        buf = export_csv(some_rows)
        text = buf.getvalue().decode("utf-8-sig")
        # All three entity_ids
        assert "cc-0" in text and "cc-1" in text and "cc-2" in text
        # Categories preserved
        assert "master_data" in text


class TestExportXLSX:
    def test_writes_xlsx_when_openpyxl_available(self, some_rows):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            pytest.skip("openpyxl not installed")
        buf = export_xlsx(some_rows, exported_by="Anna")
        data = buf.getvalue()
        # XLSX is a zip file → starts with PK\x03\x04
        assert data.startswith(b"PK")
        # Open and verify rows
        from openpyxl import load_workbook
        import io
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        # First row is title; column headers exist somewhere
        all_text = ""
        for row in ws.iter_rows(values_only=True):
            for cell in row:
                if cell is not None:
                    all_text += str(cell) + " "
        assert "id" in all_text
        assert "category" in all_text
        assert "cc-0" in all_text


class TestExportRouter:
    def test_csv_default(self, test_client, some_rows):
        resp = test_client.get("/api/audit/export", headers=HEADERS_CTRL)
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("text/csv")
        assert "audit_log_" in resp.headers.get("content-disposition", "")
        # Total header set
        assert "X-Audit-Export-Total" in resp.headers

    def test_xlsx_format(self, test_client, some_rows):
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            pytest.skip("openpyxl not installed")
        resp = test_client.get(
            "/api/audit/export?format=xlsx", headers=HEADERS_CTRL,
        )
        assert resp.status_code == 200
        # Either real xlsx OR fallback CSV — both are acceptable per design.
        ct = resp.headers["content-type"]
        if "spreadsheet" in ct:
            assert resp.content.startswith(b"PK")
        else:
            assert resp.headers.get("X-Audit-Export-Fallback") is not None

    def test_filter_passes_through(self, test_client, some_rows):
        resp = test_client.get(
            "/api/audit/export?category=master_data", headers=HEADERS_CTRL,
        )
        assert resp.status_code == 200
        # All seeded rows are master_data
        assert int(resp.headers["X-Audit-Export-Total"]) == 3
